#!/usr/bin/env python
"""Append Lindblad / self-Kerr teaching slides to the existing GDR deck.

Does not rebuild the deck (manual edits are kept).
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "report_assets" / "figures"
PPTX = HERE / "gaussian_data_regression.pptx"

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

FIGURES = {
    "lindblad_on_rho_eqs": r"""
\frac{d\rho}{dt}
  &= \sum_k \Bigl(
       c_k\rho c_k^\dagger
       -\tfrac12\{c_k^\dagger c_k,\rho\}
     \Bigr)
  \qquad (H=0\text{ during the idle}) \\[0.55em]
\rho(\tau)
  &= e^{\mathcal{L}\tau}[\rho]
  \;=\;
  \sum_j K_j\rho K_j^\dagger \\[0.65em]
c_\downarrow\rho c_\downarrow^\dagger
  &\propto a\rho a^\dagger
  \qquad\text{moves }|n\rangle\langle n|
  \;\to\;
  |n-1\rangle\langle n-1| \\[0.35em]
c_\uparrow\rho c_\uparrow^\dagger
  &\propto a^\dagger\rho a
  \qquad\text{moves }|n\rangle\langle n|
  \;\to\;
  |n+1\rangle\langle n+1| \\[0.35em]
-\tfrac12\{c^\dagger c,\rho\}
  &\qquad\text{the price of those jumps: keeps }\mathrm{Tr}\,\rho=1
""",
    "self_kerr_on_rho_eqs": r"""
H &= K\,\hat n(\hat n-1)/2,
  \qquad
  U=e^{-i H\tau},
  \qquad
  K=2\pi\times 500\,\mathrm{Hz} \\[0.45em]
U|n\rangle
  &= e^{-i\phi_n}|n\rangle,
  \qquad
  \phi_n = K\tau\, n(n-1)/2 \\[0.45em]
\langle n|U\rho U^\dagger|n\rangle
  &= \langle n|\rho|n\rangle
  \qquad\text{populations do not hop} \\[0.45em]
\langle n|U\rho U^\dagger|m\rangle
  &= e^{-i(\phi_n-\phi_m)}\langle n|\rho|m\rangle
  \qquad\text{coherences pick up a phase} \\[0.55em]
\mathcal{E}_{\mathrm{comp}}
  &= \mathcal{E}_2\circ\mathcal{E}_1\circ\mathcal{E}_q
     \circ\mathcal{U}_2\circ\mathcal{U}_1
""",
}


def _tex_document(body: str) -> str:
    return rf"""\documentclass[12pt]{{article}}
\usepackage{{fix-cm}}
\usepackage{{amsmath}}
\usepackage{{amssymb}}
\usepackage[
  paperwidth=11.8in,
  paperheight=6.4in,
  margin=0.18in,
  headheight=0pt,
  headsep=0pt,
  footskip=0pt
]{{geometry}}
\pagestyle{{empty}}
\setlength{{\abovedisplayskip}}{{2pt}}
\setlength{{\belowdisplayskip}}{{2pt}}
\begin{{document}}
\fontsize{{20}}{{26}}\selectfont
\begin{{align*}}
{body.strip()}
\end{{align*}}
\end{{document}}
"""


def _crop_white(path: Path, pad: int = 16) -> Image.Image:
    img = Image.open(path).convert("RGB")
    arr = np.asarray(img.convert("L"))
    ink = arr < 248
    if not ink.any():
        raise RuntimeError(f"blank render: {path}")
    ys, xs = np.where(ink)
    return img.crop(
        (
            max(0, int(xs.min()) - pad),
            max(0, int(ys.min()) - pad),
            min(img.width, int(xs.max()) + pad + 1),
            min(img.height, int(ys.max()) + pad + 1),
        )
    )


def render_one(name: str, body: str) -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    dest = FIG / f"{name}.png"
    pdflatex = shutil.which("pdflatex") or "/Library/TeX/texbin/pdflatex"
    qlmanage = "/usr/bin/qlmanage"
    tex_source = _tex_document(body)
    (FIG / f"{name}.tex").write_text(tex_source, encoding="utf-8")
    with tempfile.TemporaryDirectory(prefix=f"gdr_{name}_") as tmp:
        td = Path(tmp)
        tex_path = td / f"{name}.tex"
        tex_path.write_text(tex_source, encoding="utf-8")
        proc = subprocess.run(
            [
                pdflatex,
                "-interaction=nonstopmode",
                "-halt-on-error",
                f"-output-directory={td}",
                str(tex_path),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        pdf = td / f"{name}.pdf"
        if proc.returncode != 0 or not pdf.is_file():
            raise RuntimeError(proc.stdout[-2500:] + "\n" + proc.stderr[-800:])
        subprocess.run(
            [qlmanage, "-t", "-s", "3600", "-o", str(td), str(pdf)],
            check=True,
            capture_output=True,
            text=True,
        )
        img = _crop_white(td / f"{name}.pdf.png")
        left, right, tb = 10, 14, 6
        canvas = Image.new(
            "RGB",
            (max(3600, img.width + left + right), img.height + 2 * tb),
            (255, 255, 255),
        )
        canvas.paste(img, (left, tb))
        canvas.save(dest, "PNG")
        print(f"wrote {dest} {canvas.size}")
    return dest


def _run(p, text, size, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = BLACK
    return r


def _add_textbox(slide, x, y, w, h, lines):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, size, bold, after in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _run(p, text, size, bold=bold)
        p.space_after = Pt(after)
    return tf


def _white_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return s


def append_slides() -> None:
    prs = Presentation(str(PPTX))

    s = _white_slide(prs)
    _add_textbox(
        s,
        0.7,
        0.26,
        12.0,
        0.50,
        [("Lindblad, acting on the density matrix", 26, True, 0)],
    )
    _add_textbox(
        s,
        0.7,
        0.80,
        12.0,
        0.70,
        [
            (
                "During each idle of length τ the code uses H = 0. Dissipation is a first-order ODE on ρ, not a single matrix multiply.",
                16,
                False,
                4,
            ),
            (
                "The jump term c ρ c† actually moves probability. The anticommutator is only bookkeeping so that Tr ρ stays 1. The finite-time map is then written as Kraus operators and applied locally.",
                16,
                False,
                0,
            ),
        ],
    )
    png = FIG / "lindblad_on_rho_eqs.png"
    im = Image.open(png)
    w = 12.2
    h = w * im.size[1] / im.size[0]
    s.shapes.add_picture(str(png), Inches(0.55), Inches(1.62), width=Inches(w))
    y = min(1.62 + h + 0.10, 6.55)
    _add_textbox(
        s,
        0.7,
        y,
        12.0,
        0.80,
        [
            (
                "Concrete photon loss: if the cavity is in |n⟩⟨n|, the jump a ρ a† produces |n−1⟩⟨n−1| at a rate proportional to n. That is why loss is a classical hop on the histogram.",
                16,
                False,
                4,
            ),
            (
                "Transmon T1 is the same idea on the qubit: σ− ρ σ+ sends |1⟩⟨1| to |0⟩⟨0|. Number dephasing uses c = √κφ n̂: no hops in n, only extra decay of coherences ⟨n|ρ|m⟩.",
                16,
                False,
                0,
            ),
        ],
    )

    s = _white_slide(prs)
    _add_textbox(
        s,
        0.7,
        0.26,
        12.0,
        0.50,
        [("Self-Kerr, acting on the density matrix", 26, True, 0)],
    )
    _add_textbox(
        s,
        0.7,
        0.80,
        12.0,
        0.85,
        [
            (
                "Self-Kerr is not a jump. It is ordinary Hamiltonian evolution of the cavity, U ρ U†, from the photon-photon interaction in the oscillator.",
                16,
                False,
                4,
            ),
            (
                "Every Fock state |n⟩ only picks up a phase that grows like n(n−1). The number histogram does not change. Superpositions |n⟩+|m⟩ pick up a relative phase, so later non-covariant errors (transmon, control) can see a different state.",
                16,
                False,
                0,
            ),
        ],
    )
    png = FIG / "self_kerr_on_rho_eqs.png"
    im = Image.open(png)
    w = 12.2
    h = w * im.size[1] / im.size[0]
    s.shapes.add_picture(str(png), Inches(0.55), Inches(1.78), width=Inches(w))
    y = min(1.78 + h + 0.10, 6.55)
    _add_textbox(
        s,
        0.7,
        y,
        12.0,
        0.85,
        [
            (
                "Comprehensive apply order on ρ, after the (warped) gate: Kerr on cavity 1, Kerr on cavity 2, then transmon Lindblad, then cavity-1 Lindblad, then cavity-2 Lindblad. Each Lindblad map is Σ K ρ K† on that factor.",
                16,
                False,
                4,
            ),
            (
                "Readout is after p(q,n,m) = ⟨q n m|ρ|q n m⟩, never inside these maps. Cross-Kerr and residual σz(n1+n2) exist in code and are off.",
                16,
                False,
                0,
            ),
        ],
    )

    prs.save(str(PPTX))
    print(f"saved {PPTX}  slides={len(prs.slides)}")


if __name__ == "__main__":
    for name, body in FIGURES.items():
        render_one(name, body)
    append_slides()
