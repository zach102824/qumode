#!/usr/bin/env python
"""Build a plain white GDR intro deck.

Output: Error_mitigation/gaussian_data_regression.pptx
Facts from Error_mitigation_report.tex / official gdr_param matrix (108 trials).
"""

from __future__ import annotations

import importlib.util
import os

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "report_assets", "figures")
OUT = os.path.join(HERE, "gaussian_data_regression.pptx")


def _run_renderer(module_name: str, filename: str) -> None:
    path = os.path.join(HERE, "report_assets", filename)
    spec = importlib.util.spec_from_file_location(module_name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.main()


def _render_phase_covariant_eqs() -> None:
    _run_renderer("render_phase_covariant_eqs", "render_phase_covariant_eqs.py")


def _render_reshuffle_eqs() -> None:
    _run_renderer("render_reshuffle_eqs", "render_reshuffle_eqs.py")


def _render_slide_eqs() -> None:
    _run_renderer("render_slide_eqs", "render_slide_eqs.py")

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)


prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return s


def _run(p, text, size, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = BLACK
    return r


def _para(tf, text, size=16, bold=False, first=False, after=8, before=0, align=None, spacing=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if first and tf.paragraphs[0].runs:
        p = tf.add_paragraph()
    _run(p, text, size, bold=bold)
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    p.line_spacing = spacing
    if align is not None:
        p.alignment = align
    return p


def tb(s, x, y, w, h, *, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    return tf


def title(s, text):
    tf = tb(s, 0.7, 0.26, 12.0, 0.62)
    _para(tf, text, size=26, bold=True, first=True, after=0)


def place_fig(s, name, x, y, w, gap=0.08):
    """Place a figure and return the y just below it (plus gap)."""
    h = fig_h(name, w)
    pic(s, name, x, y, w=w)
    return y + h + gap


def pic(s, name, x, y, w=None, h=None):
    path = os.path.join(FIG, name)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"figure missing: {path}")
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def fig_h(name, w):
    path = os.path.join(FIG, name)
    im = PILImage.open(path)
    return w * im.size[1] / im.size[0]


def _set_cell_border(cell, color="000000", width_pt=0.75):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag in {qn("a:lnL"), qn("a:lnR"), qn("a:lnT"), qn("a:lnB")}:
            tcPr.remove(child)
    emu = str(int(width_pt * 12700))
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
        ln.set("w", emu)
        solid = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(solid, qn("a:srgbClr"), val=color)


def table(s, data, x, y, w, h, *, col_w=None, font=14, header=True):
    rows, cols = len(data), len(data[0])
    shp = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shp.table
    # Neutral grid. Default Accent styles paint a blue header.
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    if col_w is not None:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Emu(70000)
            cell.margin_right = Emu(70000)
            cell.margin_top = Emu(40000)
            cell.margin_bottom = Emu(40000)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            _run(p, val, font, bold=(header and i == 0) or (j == 0 and i > 0))
            _set_cell_border(cell)
    return tbl


# ---------------------------------------------------------------------------
# 1. Title
# ---------------------------------------------------------------------------
s = slide()
tf = tb(s, 0.7, 2.35, 12.0, 1.4)
_para(tf, "Gaussian Data Regression", size=36, bold=True, first=True, after=14)
_para(tf, "Error mitigation for a hybrid qubit–qumode circuit", size=20, after=0)

tf = tb(s, 0.7, 4.15, 11.5, 1.2)
_para(
    tf,
    "Same idea as Clifford Data Regression: learn the noise from circuits whose ideal answer is classically easy.",
    size=18,
    first=True,
    after=0,
)

# ---------------------------------------------------------------------------
# 2. Inspired by CDR
# ---------------------------------------------------------------------------
s = slide()
title(s, "Inspired by Clifford Data Regression")

tf = tb(s, 0.7, 1.05, 12.0, 1.35)
_para(tf, "The shared idea", size=18, bold=True, first=True, after=8)
_para(
    tf,
    "Both methods use a known, classically easy family to get the exact ideal output. They never need the ideal of the hard target circuit.",
    size=17,
    after=0,
)

tf = tb(s, 0.7, 2.55, 5.8, 4.3)
_para(tf, "CDR  (Czarnik et al., 2021)", size=17, bold=True, first=True, after=10)
_para(tf, "Replace most gates by Cliffords.", size=16, after=7)
_para(tf, "Those circuits are cheap to simulate (Gottesman–Knill).", size=16, after=7)
_para(tf, "Run them on the noisy device. Fit noisy → ideal.", size=16, after=7)
_para(tf, "Apply the fitted map to the target.", size=16, after=0)

tf = tb(s, 6.8, 2.55, 5.8, 4.3)
_para(tf, "GDR  (this work)", size=17, bold=True, first=True, after=10)
_para(tf, "Replace the hard bosonic gates by Gaussians.", size=16, after=7)
_para(tf, "Those circuits are cheap to simulate (means + covariance, or a product coherent state).", size=16, after=7)
_para(tf, "Run the same noisy channel on the twins. Fit noisy → ideal.", size=16, after=7)
_para(tf, "Apply the fitted map to the target histogram.", size=16, after=0)

# ---------------------------------------------------------------------------
# 3. One-to-one map
# ---------------------------------------------------------------------------
s = slide()
title(s, "CDR and GDR, one to one")

table(
    s,
    [
        ["", "CDR", "GDR"],
        ["Easy gate set", "Clifford", "Gaussian"],
        ["Hard leftover", "T, CCZ, …", "generic SNAP, ECD, Kerr"],
        ["Noise analog", "global depolarizing", "phase-covariant (loss, thermal, n-dephasing)"],
        ["What is fitted", "linear map on a scalar observable", "transfer matrix on the joint histogram"],
        ["Classical cost", "poly in Cliffords, exp in non-Cliffords", "poly in Gaussians, exp in non-Gaussians"],
    ],
    0.7,
    1.00,
    12.0,
    3.20,
    col_w=[2.3, 4.7, 5.0],
    font=14,
)

tf = tb(s, 0.7, 4.35, 12.0, 2.85)
_para(tf, "Why depolarizing and phase-covariant are the matching pair", size=16, bold=True, first=True, after=6)
_para(
    tf,
    "Global depolarizing collapses every scalar observable to (1 − p) × ideal + p × trash. A two-parameter fit is enough.",
    size=15,
    after=5,
)
_para(
    tf,
    "Phase-covariant noise is the matching analog for the measured histogram: one transfer matrix M, not an arbitrary map on the state. Details on the next slide.",
    size=15,
    after=0,
)

# ---------------------------------------------------------------------------
# 3b. What phase-covariant actually means  (name + positive example)
# ---------------------------------------------------------------------------
_render_phase_covariant_eqs()
s = slide()
title(s, "What phase-covariant actually means")

tf = tb(s, 0.7, 0.96, 12.0, 0.95)
_para(tf, "The name", size=17, bold=True, first=True, after=6)
_para(
    tf,
    "Phase-covariant means the noise does not pick a preferred cavity phase. Rotating the oscillator before the noise is the same as rotating after.",
    size=16,
    after=0,
)
name_y = place_fig(s, "phase_covariant_eqs.png", 0.65, 1.92, 12.0, gap=0.14)

tf = tb(s, 0.7, name_y, 12.0, 0.88)
_para(tf, "Positive example — one M exists", size=17, bold=True, first=True, after=6)
_para(
    tf,
    "Photon loss (also thermal loss, number dephasing). Take two states with the same number histogram but opposite coherence.",
    size=16,
    after=0,
)
pos_y = place_fig(s, "phase_covariant_pos_eqs.png", 0.65, name_y + 0.90, 12.0, gap=0.10)

tf = tb(s, 0.7, pos_y, 6.6, max(1.6, 7.35 - pos_y))
_para(
    tf,
    "Both have p0 = p2 = 1/2. Under loss, the sign of that coherence never enters the photon-number bins, so both produce the same q = Mp.",
    size=16,
    first=True,
    after=8,
)
_para(
    tf,
    "One column-stochastic M works for every state (end-of-circuit). That is why GDR can learn M on Gaussian twins and apply it to the target.",
    size=16,
    after=0,
)
pic(s, "phase_covariant_pos_hist.png", 7.45, pos_y, w=5.25)

# ---------------------------------------------------------------------------
# 3c. When a single M cannot exist  (negative example)
# ---------------------------------------------------------------------------
s = slide()
title(s, "When a single M cannot exist")

tf = tb(s, 0.7, 0.90, 12.0, 0.78)
_para(tf, "How D(α) treats |+⟩ and |−⟩ differently", size=17, bold=True, first=True, after=4)
_para(
    tf,
    "Two paths into n = 1: create a photon from |0⟩, or destroy one from |2⟩. The sign in |±⟩ decides whether those amplitudes cancel or add.",
    size=16,
    after=0,
)
neg_y = place_fig(s, "phase_covariant_neg_eqs.png", 0.45, 1.70, 12.4, gap=0.08)

tf = tb(s, 0.7, neg_y, 6.9, max(2.0, 7.35 - neg_y))
_para(
    tf,
    "Same p, but the n = 1 probabilities differ, so q+ ≠ q−. No single M can map that p to both histograms. A fit on twins need not apply to a target with different coherences.",
    size=15,
    first=True,
    after=8,
)
_para(
    tf,
    "A squeeze is the same kind of failure. Comprehensive noise (transmon T1/T2, ECD control error) is this on the hybrid space: M is only an effective fit.",
    size=15,
    after=0,
)
pic(s, "phase_covariant_neg_hist.png", 7.7, neg_y, w=5.2)

# ---------------------------------------------------------------------------
# 3d. Classical reshuffling, in numbers
# ---------------------------------------------------------------------------
_render_reshuffle_eqs()
_render_slide_eqs()
s = slide()
title(s, "Classical reshuffling, in numbers")

eq_w, eq_y = 12.0, 0.88
eq_h = fig_h("reshuffle_eqs.png", eq_w)
pic(s, "reshuffle_eqs.png", 0.65, eq_y, w=eq_w)

cap_y = eq_y + eq_h + 0.06
tf = tb(s, 0.7, cap_y, 12.0, 1.18)
_para(tf, "Same loss channel as the plus / minus example. One mode first. Applying M is a Markov hop of probability mass. No amplitudes.", size=15, bold=True, first=True, after=2)
_para(
    tf,
    "Ideal Fock-2 (1) puts all mass on n = 2. Pure loss with transmission 0.8 hops that column as in (2) and (4): stay 2 with 0.64, drop to 1 with 0.32, drop to 0 with 0.04. Those three numbers are exactly column 2 of M, and (3) is that column.",
    size=15,
    after=0,
)
coh_head = cap_y + 1.20
tf = tb(s, 0.7, coh_head, 12.0, 0.32)
_para(tf, "Why coherences do not matter for the histogram", size=15, bold=True, first=True, after=0)
coh_y = place_fig(s, "reshuffle_coherence_eqs.png", 0.65, coh_head + 0.32, 12.0, gap=0.05)
tf = tb(s, 0.7, coh_y, 12.0, max(1.0, 7.35 - coh_y))
_para(
    tf,
    "After the same loss, the same column-stochastic M is applied to that histogram. There is no extra interference term in the photon-number bins. That is classical reshuffling.",
    size=15,
    first=True,
    after=6,
)
_para(
    tf,
    "The real GDR matrix is the same idea on 128 bins (Kronecker product over qubit and two cavities). A non-covariant channel (a displacement, or transmon errors that mix the qubit into the cavity) can move population using coherences; then one M learned from Gaussian twins need not apply to the target.",
    size=15,
    after=0,
)

# ---------------------------------------------------------------------------
# 4. Gaussian gates + ECD / SNAP
# ---------------------------------------------------------------------------
s = slide()
title(s, "Gaussian gates, and how ECD / SNAP become them")

tf = tb(s, 0.7, 0.98, 12.0, 2.05)
_para(tf, "What a Gaussian gate is, and why it is cheap", size=17, bold=True, first=True, after=8)
_para(
    tf,
    "Generated by a Hamiltonian at most quadratic in q, p. Common gates: displacement, rotation, squeezing, beam splitter.",
    size=16,
    after=6,
)
_para(
    tf,
    "A Gaussian state is fixed by a 2m-vector of means and a 2m by 2m covariance. Updates are polynomial in the number of modes, independent of any Fock cutoff.",
    size=16,
    after=6,
)
_para(
    tf,
    "Cost of the leftover hard gates is exponential in how many you keep (Gaussian rank / cutoff). Same split as Clifford vs T.",
    size=16,
    after=0,
)

tf = tb(s, 0.7, 3.12, 5.9, 1.15)
_para(tf, "ECD to Gaussian", size=17, bold=True, first=True, after=6)
_para(tf, "Snap each ancilla angle onto {0, pi}.", size=16, after=4)
_para(tf, "The ancilla stays in the computational basis.", size=16, after=0)
ecd_y = place_fig(s, "ecd_gaussian_eqs.png", 0.70, 4.28, 5.9)
tf = tb(s, 0.7, ecd_y, 5.9, max(0.8, 7.35 - ecd_y))
_para(
    tf,
    "From vacuum that is the state. The ideal histogram is the truncated one-mode product evolution (approximately Poisson).",
    size=16,
    first=True,
    after=0,
)

tf = tb(s, 6.8, 3.12, 5.9, 1.15)
_para(tf, "SNAP to Gaussian", size=17, bold=True, first=True, after=6)
_para(tf, "Replace each phase list by its least-squares affine fit (constant term zero).", size=16, after=0)
snap_y = place_fig(s, "snap_gaussian_eqs.png", 6.80, 4.28, 5.9)
tf = tb(s, 6.8, snap_y, 5.9, max(0.8, 7.35 - snap_y))
_para(tf, "What remains is displacements and rotations.", size=16, first=True, after=6)
_para(
    tf,
    "A few original SNAP / ECD gates can be left in (t_free = 2) so the twins sit closer to the target.",
    size=16,
    after=0,
)

# ---------------------------------------------------------------------------
# 5. Method steps
# ---------------------------------------------------------------------------
s = slide()
title(s, "How GDR is run")

gdr_y = place_fig(s, "gdr_run_eqs.png", 0.65, 0.92, 12.0)
steps = [
    ("1. Target", [
        "ECD depth 5 or SNAP depth 2. Score the joint q, n, m histogram against the ideal.",
    ]),
    ("2. Twins", [
        "40 twins, same depth and gate count as the target. Mix is 75% / 25% of the twins, not 25% of the gates in each circuit.",
        "75% (30) fully Gaussian (t_free = 0): every ECD angle snapped to {0, pi}; every SNAP phase list replaced by an affine fit.",
        "25% (10) keep t_free = 2 original hard gates, drawn uniformly at random among the (layer, cavity) slots: two of 10 ECD–rotation pairs, or two of 4 SNAP lists.",
        "ECD leftover: unsapped angles, so those ECDs stay the original conditional ECD, not the unconditional displacement plus bit-flip. SNAP leftover: original non-affine phases, not the affine fit. All other gates in those twins are still Gaussianized. Gaussian rank is 2 to the power t_free.",
    ]),
    ("3. Ideal twins", [
        "Classical. Fully Gaussian twins use product one-mode evolution, checked against the statevector (TVD below 1e-6).",
    ]),
    ("4. Same noise", [
        "Target and twins see the same circuit-noise family, the same readout, and 8,192 shots.",
    ]),
    ("5. Fit M", [
        "11 bounded parameters (Kronecker M above), multinomial MLE on the twins.",
    ]),
    ("6. Unfold", [
        "Richardson–Lucy on the observed histogram. Keeps probabilities non-negative and normalized. This is gdr_param.",
    ]),
]
tf = tb(s, 0.7, gdr_y, 12.0, max(3.5, 7.35 - gdr_y))
first = True
for head, bodies in steps:
    _para(tf, head, size=16, bold=True, first=first, after=1, before=0 if first else 4)
    for body in bodies:
        _para(tf, body, size=15, after=1)
    first = False

# ---------------------------------------------------------------------------
# 6. How circuit noise hits ρ
# ---------------------------------------------------------------------------
s = slide()
title(s, "Circuit noise is applied to the density matrix")

tf = tb(s, 0.7, 0.90, 12.0, 0.58)
_para(
    tf,
    "Hybrid space: qubit tensor two 8-level cavities, so the density matrix is 128 by 128. Circuit noise is a CPTP map on that state, interleaved with the gates. Readout is not in the channel — it is a confusion matrix on the final q, n, m probabilities.",
    size=15,
    first=True,
    after=0,
)
apply_y = place_fig(s, "circuit_noise_apply_eqs.png", 0.65, 1.50, 12.0, gap=0.04)
tf = tb(s, 0.7, apply_y, 12.0, 0.26)
_para(tf, "That histogram is what readout and shots see.", size=15, first=True, after=0)

table_y = apply_y + 0.28
table(
    s,
    [
        ["Family", "When the channel is applied", "Applications (ECD N_d=5 / SNAP N_d=2)"],
        ["Loss", "after each UER / SNAP layer", "5 / 2"],
        ["Loss + thermal + dephasing", "after each UER / SNAP layer", "5 / 2"],
        ["Comprehensive", "after each ECD–rotation or SNAP–D pair", "10 / 4"],
    ],
    0.7,
    table_y,
    12.0,
    1.52,
    col_w=[3.4, 5.0, 3.6],
    font=13,
)

ord_y = table_y + 1.58
tf = tb(s, 0.7, ord_y, 12.0, 0.90)
_para(tf, "One call of the channel, in this order", size=15, bold=True, first=True, after=3)
_para(tf, "1. Kerr unitary on cavity 1, then cavity 2          (comprehensive only; not a Kraus map)", size=14, after=2)
_para(tf, "2. Transmon Kraus on the qubit                     (comprehensive only)", size=14, after=2)
_para(tf, "3. Cavity-1 Kraus, then cavity-2 Kraus             (every family; the two modes are independent)", size=14, after=0)
kraus_y = place_fig(s, "circuit_noise_kraus_eqs.png", 0.65, ord_y + 0.92, 12.0, gap=0.04)
tf = tb(s, 0.7, kraus_y, 12.0, max(0.55, 7.40 - kraus_y))
_para(
    tf,
    "The 16,384 by 16,384 Liouvillian is never built. Same kτ each time, so comprehensive ECD sees twice as many loss maps. Loss and thermal stay phase-covariant; comprehensive (transmon, Kerr, control) does not — then a histogram M is only an effective fit.",
    size=15,
    first=True,
    after=0,
)

# ---------------------------------------------------------------------------
# 6b. Loss family Kraus
# ---------------------------------------------------------------------------
s = slide()
title(s, "Loss family: paper amplitude-damping Kraus")

tf = tb(s, 0.7, 0.95, 12.0, 0.85)
_para(
    tf,
    "Pure photon loss, n_th = 0. The same Kraus set is applied independently to each cavity. No qubit channel. No Kerr. This is the only Kraus set in the loss family.",
    size=16,
    first=True,
    after=0,
)

tf = tb(s, 0.7, 1.85, 12.0, 0.38)
_para(tf, "Operators on one cavity (cutoff L = 8)", size=17, bold=True, first=True, after=0)
pic(s, "kraus_operators.png", 0.70, 2.22, h=2.22)

tf = tb(s, 0.7, 4.50, 12.0, 0.72)
_para(tf, "What each jump operator does", size=16, bold=True, first=True, after=4)
_para(
    tf,
    "The j-th operator lowers the photon number by exactly j. The first is n to n-1, the second is n to n-2, and so on. Nothing in this set can raise n. The no-jump operator is not quite the infinite-dimension form; on a truncated L-level space it is the residual square root so completeness holds.",
    size=15,
    after=0,
)
cap_y = place_fig(s, "kraus_caption_eqs.png", 0.65, 5.24, 12.0, gap=0.06)
tf = tb(s, 0.7, cap_y, 12.0, max(0.6, 7.40 - cap_y))
_para(
    tf,
    "Loss is not measurement error. The jump probability grows with n (binomial in the transmission). A detector flip p_nn does not. That distinction is why the two are fitted as separate parameters later.",
    size=15,
    first=True,
    after=0,
)

# ---------------------------------------------------------------------------
# 6c. Lindblad Kraus
# ---------------------------------------------------------------------------
s = slide()
title(s, "Thermal and comprehensive: Lindblad collapse operators, then Kraus")

tf = tb(s, 0.7, 0.90, 12.0, 0.62)
_para(
    tf,
    "These two families do not use the paper amplitude-damping Kraus set. On each local space they take collapse operators, build the H = 0 Liouvillian, and convert the finite-time map to a Kraus set. Those Kraus operators are what the channel applies.",
    size=15,
    first=True,
    after=0,
)
intro_y = place_fig(s, "lindblad_intro_eqs.png", 0.65, 1.52, 12.0, gap=0.06)

tf = tb(s, 0.7, intro_y, 6.05, 0.36)
_para(tf, "Cavity collapse operators (both cavities, independently)", size=15, bold=True, first=True, after=0)
tf = tb(s, 6.90, intro_y, 5.85, 0.36)
_para(
    tf,
    "Transmon collapse operators (comprehensive only; T1 = 50 us, T2 = 30 us)",
    size=15,
    bold=True,
    first=True,
    after=0,
)
cav_y = place_fig(s, "lindblad_cavity_eqs.png", 0.62, intro_y + 0.36, 6.10, gap=0.04)
tr_y = place_fig(s, "lindblad_transmon_eqs.png", 6.85, intro_y + 0.36, 5.90, gap=0.04)
ops_y = max(cav_y, tr_y)

tf = tb(s, 0.7, ops_y, 12.0, 0.42)
_para(
    tf,
    "The collapse rate is chosen so that kappa times the idle time equals the listed kτ for one application. Tau is the idle time of that application. Qubit thermal occupation is 0, so there is no raising-operator heating term.",
    size=14,
    first=True,
    after=0,
)

table(
    s,
    [
        ["", "n_th", "k_phi tau", "qubit Kraus", "Kerr / control"],
        ["Loss + thermal + dephasing", "0.05", "0.5 kτ", "off", "off"],
        ["Comprehensive", "0.01", "0", "on", "on"],
    ],
    0.7,
    ops_y + 0.46,
    12.0,
    1.38,
    col_w=[3.5, 1.5, 1.8, 2.4, 2.8],
    font=14,
)

# ---------------------------------------------------------------------------
# 6d. Not Kraus + readout
# ---------------------------------------------------------------------------
s = slide()
title(s, "On the density matrix but not Kraus, and readout after")

tf = tb(s, 0.7, 0.90, 12.0, 0.70)
_para(tf, "Comprehensive extras that are not a Kraus map", size=17, bold=True, first=True, after=4)
_para(
    tf,
    "Self-Kerr, applied before the Kraus maps, independently on each cavity:",
    size=16,
    after=0,
)
kerr_y = place_fig(s, "kerr_eqs.png", 0.65, 1.60, 12.0, gap=0.05)
tf = tb(s, 0.7, kerr_y, 12.0, 0.55)
_para(
    tf,
    "Control error is not a channel at all. The compiled ECD parameters are replaced once (SNAP compiled parameters are not warped). That changes the circuit, not a stochastic kernel on the histogram.",
    size=16,
    first=True,
    after=0,
)
ctrl_y = place_fig(s, "control_error_eqs.png", 0.65, kerr_y + 0.55, 12.0, gap=0.08)

tf = tb(s, 0.7, ctrl_y, 12.0, 0.36)
_para(tf, "Readout: after the density matrix is finished, never on the state", size=17, bold=True, first=True, after=0)
pobs_y = place_fig(s, "readout_pobs_eqs.png", 0.65, ctrl_y + 0.36, 12.0, gap=0.04)
tf = tb(s, 0.7, pobs_y, 12.0, 1.15)
_para(
    tf,
    "C_q is the 2 by 2 bit-flip matrix with rates (p01, p10). Each Fock matrix sends n to n+1 or n-1 with p_nn/2 each way (total hop p_nn), and stays with 1 - p_nn.",
    size=15,
    first=True,
    after=4,
)
_para(tf, "ideal: (0, 0) and p_nn = 0", size=15, after=2)
_para(tf, "realistic: (0.01, 0.03) and p_nn = 0.03", size=15, after=2)
_para(tf, "strong: (0.03, 0.08) and p_nn = 0.10", size=15, after=0)
zne_y = place_fig(s, "zne_scale_eqs.png", 0.65, pobs_y + 1.18, 12.0, gap=0.04)
tf = tb(s, 0.7, zne_y, 12.0, max(0.45, 7.40 - zne_y))
_para(
    tf,
    "One mixed p-spin instance (H000). Idle-time ZNE does not scale this detector or n_th.",
    size=15,
    first=True,
    after=0,
)

# ---------------------------------------------------------------------------
# 7. Transfer matrix
# ---------------------------------------------------------------------------
s = slide()
title(s, "The transfer matrix: 11 parameters, not 16,384")

tm_y = place_fig(s, "transfer_matrix_eqs.png", 0.65, 0.90, 12.0, gap=0.06)
tf = tb(s, 0.7, tm_y, 12.0, 0.72)
_para(
    tf,
    "Histogram is 2 by 8 by 8 = 128 bins. A free 128 by 128 matrix has 16,384 entries. Even a free Kronecker model is 4 + 64 + 64 = 132. GDR keeps a physics-shaped kernel with 11 numbers. The suffix 1 vs 2 is cavity 1 (photon number n, factor C1) vs cavity 2 (m, factor C2).",
    size=15,
    first=True,
    after=0,
)

table(
    s,
    [
        ["#", "Parameter", "What it does to M"],
        ["1–2", "eta_1, eta_2", "Per cavity. Transmission in that mode’s 8 by 8 thermal-loss kernel. With n_th = 0, only lowers n. Circuit-like."],
        ["3–4", "n_th1, n_th2", "Per cavity. Mean thermal photons in that mode’s bath. Opens n to n+1 heating in the same kernel. Not the state’s mean photon number, not readout."],
        ["5–7", "p_down, p_up, eps", "Shared by both cavities. Extra n-independent hops n to n-1 / n to n+1, then leak toward a uniform Fock row. Circuit residual."],
        ["8–9", "p01, p10", "Only Cq. Qubit readout: P(measure 1 | 0) and P(measure 0 | 1). Moves q = 0 or 1. Leaves n and m alone."],
        ["10–11", "p_nn1, p_nn2", "Per cavity, applied last. Detector n to n+1 or n-1 at p_nn/2 each way. Rate does not scale with n. Readout-like."],
    ],
    0.7,
    tm_y + 0.78,
    12.0,
    3.10,
    col_w=[1.2, 2.8, 8.0],
    font=13,
)

tf = tb(s, 0.7, tm_y + 3.96, 12.0, 1.55)
_para(
    tf,
    "Each C_mode is built in order: thermal-loss(eta, n_th), then p_down, then p_up, then eps, then p_nn. Cq is only the 2 by 2 bit-flip. Fit the 11 numbers on the twins by multinomial MLE, then invert M on the target with Richardson–Lucy (80 iterations).",
    size=15,
    first=True,
    after=5,
)
_para(
    tf,
    "Loss and readout are fitted jointly: n to n-1 from eta and n to n+1 or n-1 from p_nn look the same on one circuit. n_th is the other way a cavity factor can raise n — next slide. The transfer-matrix relation is exact for end-of-circuit phase-covariant noise, or a fully Gaussian twin; interleaved leftover non-Gaussian gates, and comprehensive, make M an effective fit.",
    size=15,
    after=0,
)

# ---------------------------------------------------------------------------
# 8. What each of the 11 parameters does  (immediately after the 11-param table)
# ---------------------------------------------------------------------------
s = slide()
title(s, "What each parameter does to M")

tf = tb(s, 0.7, 0.90, 12.0, 0.58)
_para(
    tf,
    "C1 is cavity 1 (n), C2 is cavity 2 (m), each 8 by 8. Build order: thermal-loss(eta, n_th), then p_down, then p_up, then eps, then p_nn. p_down, p_up, eps are shared — the same three numbers go into both cavities. eta, n_th, and p_nn are per cavity.",
    size=15,
    first=True,
    after=0,
)

tf = tb(s, 0.7, 1.50, 6.0, 1.35)
_para(tf, "Per-mode loss / thermal  (circuit-like)", size=16, bold=True, first=True, after=4)
_para(
    tf,
    "eta_1, eta_2 — Transmission of that cavity. The 8 by 8 population map is the matrix exponential of a truncated birth-death generator. eta = 1 is no loss. If n_th = 0 the map is binomial: mass only stays or moves to lower n. eta cannot heat by itself.",
    size=13,
    after=0,
)
bd_y = place_fig(s, "birth_death_eqs.png", 0.70, 2.88, 6.0, gap=0.06)
tf = tb(s, 0.7, bd_y, 6.0, max(2.4, 7.40 - bd_y))
_para(
    tf,
    "n_th1, n_th2 — Mean thermal photons in that cavity’s bath (how warm the loss port is). Not the photon number of the quantum state, and not a second readout error. The up-rate heats (n to n+1) in that mode’s factor of M. They also multiply the down-rate. If n_th = 0, that factor cannot raise n. If eta = 1, the rates vanish and n_th does nothing.",
    size=13,
    first=True,
    after=8,
)
_para(tf, "Shared hops / leak  (circuit residual; both cavities)", size=16, bold=True, after=4)
_para(
    tf,
    "p_down — After loss, hop n to n-1 with probability p_down, independent of n (no hop at n = 0). Cannot raise n.",
    size=13,
    after=4,
)
_para(
    tf,
    "p_up — Hop n to n+1 with probability p_up, independent of n. Can raise n, but the rate does not grow with (n+1) the way thermal heating does.",
    size=13,
    after=4,
)
_para(
    tf,
    "eps — Mix a fraction eps of each column toward the uniform 1/8 Fock distribution. Mass from any n can land on any bin, including higher n.",
    size=13,
    after=0,
)

tf = tb(s, 6.9, 1.50, 5.8, 5.70)
_para(tf, "Qubit readout  (readout-like; only Cq)", size=16, bold=True, first=True, after=6)
_para(
    tf,
    "p01 = P(measure 1 | true 0).  p10 = P(measure 0 | true 1). Moves probability between the q = 0 and q = 1 slices. Does not change n or m.",
    size=13,
    after=10,
)
_para(tf, "Fock readout  (readout-like; last on that cavity)", size=16, bold=True, after=6)
_para(
    tf,
    "p_nn1, p_nn2 — Detector misassignment: true n stays with 1 − p_nn, or is reported as n+1 or n-1 with p_nn/2 each way (edges keep the leftover). Can raise or lower the recorded n by one. The rate is flat in n — that is how twins separate it from eta, whose hop probability grows with n.",
    size=13,
    after=10,
)
_para(tf, "eta vs n_th vs p_nn", size=16, bold=True, after=6)
_para(
    tf,
    "eta with n_th = 0 only lowers n, and more so at large n.",
    size=13,
    after=5,
)
_para(
    tf,
    "n_th is bath occupation: it is what lets the same thermal-loss kernel also heat (n to n+1). It is not a readout parameter and not the state’s mean photon number.",
    size=13,
    after=5,
)
_para(
    tf,
    "p_nn is a detector mix-up after the circuit, n to n+1 or n-1 at a rate that does not scale like thermal loss.",
    size=13,
    after=0,
)

# ---------------------------------------------------------------------------
# 9. Results
# ---------------------------------------------------------------------------
s = slide()
title(s, "How well it works  (official gdr_param, 108 trials)")

tvd_y = place_fig(s, "tvd_eqs.png", 0.70, 0.92, 5.4, gap=0.06)
tf = tb(s, 0.7, tvd_y, 5.4, max(4.5, 7.35 - tvd_y))
_para(
    tf,
    "Total variation distance between the ideal and (raw or mitigated) q, n, m histograms. 0 is a perfect match; 1 is disjoint support. Histogram distance, not energy.",
    size=13,
    first=True,
    after=6,
)
_para(tf, "What the 108 trials are", size=15, bold=True, after=3)
_para(
    tf,
    "2 ansatze (ECD, SNAP)  ×  2 circuits (random / optimized)  ×  3 noise families  ×  3 κτ  ×  3 readout levels.",
    size=13,
    after=2,
)
_para(
    tf,
    "Families: loss, thermal+dephasing, comprehensive.  κτ ∈ {0.003, 0.03, 0.1}.  Readout: ideal, realistic, strong.",
    size=13,
    after=2,
)
_para(tf, "One instance (H000). Every trial: 8,192 shots and 40 twins.", size=13, after=8)
_para(tf, "Mean TVD", size=15, bold=True, after=1)
_para(tf, "0.331  →  0.185     (44% down)", size=17, after=5)
_para(tf, "Median TVD", size=15, bold=True, after=1)
_para(tf, "0.285  →  0.127", size=17, after=5)
_para(tf, "TVD improves in 93 / 108 trials", size=15, after=2)
_para(tf, "Optimized circuits: 54 / 54", size=15, after=5)
_para(tf, "Mean-TVD drop by family", size=14, bold=True, after=2)
_para(tf, "loss 49%   ·   thermal 46%   ·   comprehensive 39%", size=14, after=4)
_para(tf, "by κτ", size=14, bold=True, after=2)
_para(tf, "0.003 → 65%    0.03 → 57%    0.1 → 33%", size=14, after=4)
_para(tf, "The 15 regressions are almost all ECD / random at κτ ≥ 0.03 (over-correction).", size=13, after=0)

pic(s, "paired_tvd.png", 6.2, 1.05, w=6.5)

# ---------------------------------------------------------------------------
# 10. Example SNAP (official H000)
# ---------------------------------------------------------------------------
s = slide()
title(s, "Example: SNAP, mid noise, strong readout error")

tf = tb(s, 0.7, 0.95, 12.0, 0.62)
_para(
    tf,
    "One official H000 trial. Optimized SNAP · loss + thermal + dephasing · κτ = 0.03. Strong readout: qubit flips (p01, p10) = (0.03, 0.08) and Fock n to n+1 or n-1 total hop 0.10.  TVD  0.469 → 0.099.   energy error  1.72 → 0.52.",
    size=15,
    first=True,
    after=0,
)
pic(s, "case_histogram_snap.png", 0.55, 1.65, w=12.2)

# ---------------------------------------------------------------------------
# 11. Example official ECD |070⟩ (official H000, matching SNAP trial)
# ---------------------------------------------------------------------------
s = slide()
title(s, "Example: official ECD, mid noise, strong readout error")

tf = tb(s, 0.7, 0.95, 12.0, 0.62)
_para(
    tf,
    "Official H000 trial, same noise as the SNAP trial. Optimized ECD · loss + thermal + dephasing · κτ = 0.03 · strong readout. Ideal peak is ket 070.  TVD  0.445 → 0.105.   energy error  1.62 → 0.54.",
    size=15,
    first=True,
    after=0,
)
pic(s, "case_histogram_ecd.png", 0.55, 1.65, w=12.2)

# ---------------------------------------------------------------------------
# 12. Example ECD, fully Gaussian twins, optimized |016>
# ---------------------------------------------------------------------------
s = slide()
title(s, "Example: optimized ECD on ket 016, fully Gaussian twins")

tf = tb(s, 0.7, 0.95, 12.0, 0.62)
_para(
    tf,
    "Research replay on H004 (not the official 108-trial H000 matrix). Ideal peak is the true GS ket 016 (p about 0.93). Comprehensive, κτ = 0.003, same strong readout. All 40 twins fully Gaussian.  TVD  0.311 → 0.115.   energy error  1.38 → 0.39.",
    size=15,
    first=True,
    after=0,
)
pic(s, "case_histogram_ecd_gaussian.png", 0.55, 1.65, w=12.2)

# ---------------------------------------------------------------------------
# 13. Example ECD, official mix, optimized |016>
# ---------------------------------------------------------------------------
s = slide()
title(s, "Example: optimized ECD on ket 016, official twin mix")

tf = tb(s, 0.7, 0.95, 12.0, 0.62)
_para(
    tf,
    "Same H004 ket-016 target and strong readout. Official twin mix: 30 fully Gaussian plus 10 that keep two leftover ECD gates.  TVD  0.311 → 0.094.   energy error  1.38 → 0.25.",
    size=15,
    first=True,
    after=0,
)
pic(s, "case_histogram_ecd_016.png", 0.55, 1.65, w=12.2)

# ---------------------------------------------------------------------------
# 14. Limits
# ---------------------------------------------------------------------------
s = slide()
title(s, "TVD success is not energy success")

tf = tb(s, 0.7, 1.0, 5.5, 5.9)
_para(tf, "Mean energy error   0.935 → 0.488", size=17, first=True, after=6)
_para(tf, "Median energy error  0.233 → 0.279   (goes up)", size=17, after=10)
_para(tf, "Energy improves in 64 / 108 trials, worsens in 44.", size=16, after=14)
_para(tf, "Why that is allowed", size=16, bold=True, after=6)
_para(tf, "TVD weights every histogram bin equally. Energy weights bins by the Hamiltonian. Finite-shot unfold can fix the shape and still move the energy the wrong way.", size=16, after=14)
_para(tf, "Also do not claim", size=16, bold=True, after=6)
_para(tf, "Gaussian + arbitrary photon counting is always easy (Gaussian boson sampling).", size=16, after=6)
_para(tf, "One frozen instance, one shot seed. The 93 / 108 and 15 / 108 counts are that draw, not a confidence interval.", size=16, after=0)

pic(s, "energy_tradeoff.png", 6.3, 1.05, w=6.4)

prs.save(OUT)
print(OUT)
